import os
import io
from typing import List, Tuple

import fitz
from docx import Document as DocxDocument
from pptx import Presentation

try:
    from PIL import Image
    import pytesseract
    OCR_AVAILABLE = True
except Exception:
    OCR_AVAILABLE = False


def _extract_pdf(data: bytes) -> Tuple[str, int]:
    pages = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            pages.append(f"\n--- Page {i} ---\n" + page.get_text())
    return "\n".join(pages), len(pages)


def _extract_docx(data: bytes) -> Tuple[str, int]:
    doc = DocxDocument(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    return text, 1


def _extract_pptx(data: bytes) -> Tuple[str, int]:
    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_text = [f"\n--- Slide {i} ---"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text)
        slides.append("\n".join(slide_text))
    return "\n".join(slides), len(prs.slides)


def _extract_text(data: bytes) -> Tuple[str, int]:
    return data.decode("utf-8", errors="ignore"), 1


def _extract_image(data: bytes) -> Tuple[str, int]:
    if not OCR_AVAILABLE:
        return "[OCR not available]", 1
    img = Image.open(io.BytesIO(data))
    text = pytesseract.image_to_string(img)
    return text, 1


def extract_text(filename: str, data: bytes) -> Tuple[str, int]:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext in (".docx", ".doc"):
        return _extract_docx(data)
    if ext in (".pptx", ".ppt"):
        return _extract_pptx(data)
    if ext in (".txt", ".md"):
        return _extract_text(data)
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        return _extract_image(data)
    return _extract_text(data)


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> List[str]:
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: List[str] = []
    current = ""
    for p in paragraphs:
        if len(current) + len(p) + 2 <= chunk_size:
            current += ("\n\n" if current else "") + p
        else:
            if current:
                chunks.append(current)
            while len(p) > chunk_size:
                chunks.append(p[:chunk_size])
                p = p[chunk_size - overlap:]
            current = p
    if current:
        chunks.append(current)
    return chunks


def detect_metadata(filename: str, text: str) -> dict:
    lower = text.lower()
    subject_map = {
        "Networking": ["vlan", "router", "switch", "ethernet", "subnet", "ospf", "bgp", "tcp", "ip address"],
        "Cybersecurity": ["encryption", "firewall", "malware", "phishing", "vulnerability", "cve"],
        "Python": ["def ", "import ", "python", "pip", "django", "flask"],
        "Databases": ["sql", "select ", "postgres", "mysql", "table", "index", "join"],
        "Mathematics": ["theorem", "equation", "integral", "derivative", "matrix"],
        "Biology": ["cell", "dna", "protein", "organism", "enzyme"],
        "Physics": ["force", "velocity", "acceleration", "quantum", "energy"],
        "History": ["century", "war", "empire", "revolution", "treaty"],
    }
    scores = {subj: sum(lower.count(k) for k in kws) for subj, kws in subject_map.items()}
    subject = max(scores, key=scores.get) if any(scores.values()) else "General"

    topics = []
    keywords = {
        "VLANs": ["vlan", "trunk", "802.1q"],
        "Subnetting": ["subnet", "cidr", "netmask"],
        "Routing": ["router", "ospf", "bgp", "rip"],
        "Switching": ["switch", "mac address", "frame"],
        "Ethernet": ["ethernet", "csma", "frame"],
        "Security": ["firewall", "encryption", "tls", "vpn"],
        "Databases": ["sql", "join", "index"],
        "Algorithms": ["algorithm", "complexity", "sort"],
        "Python": ["def ", "class ", "import "],
    }
    for topic, kws in keywords.items():
        if any(k in lower for k in kws):
            topics.append(topic)
    if not topics:
        topics = [subject]

    difficulty = "intermediate"
    words = text.split()
    avg_word_len = sum(len(w) for w in words) / max(len(words), 1)
    if avg_word_len < 4.5:
        difficulty = "beginner"
    elif avg_word_len > 5.5:
        difficulty = "advanced"

    return {"subject": subject, "topics": topics[:6], "difficulty": difficulty}
